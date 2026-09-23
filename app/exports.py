"""
Módulo de exportación de tareas a PDF, Excel y TXT.
"""
from io import BytesIO
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ====================================================
# COLORES
# ====================================================
COLOR_DARK = "2C3E50"
COLOR_ALTA = "E74C3C"
COLOR_MEDIA = "F39C12"
COLOR_BAJA = "2ECC71"
COLOR_GRAY = "95A5A6"


# ====================================================
# PDF
# ====================================================
def generar_pdf_tareas(tareas, usuario_email):
    """Genera un PDF con todas las tareas del usuario. Devuelve bytes."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    estilos = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle(
        "TituloPersonalizado",
        parent=estilos["Title"],
        textColor=colors.HexColor("#" + COLOR_DARK),
        fontSize=20,
    )
    subtitulo_estilo = ParagraphStyle(
        "Sub",
        parent=estilos["Normal"],
        textColor=colors.HexColor("#" + COLOR_GRAY),
        fontSize=9,
        alignment=1,
    )

    elementos = []
    elementos.append(Paragraph("📋 Mis Tareas", titulo_estilo))
    elementos.append(Paragraph(
        f"Usuario: {usuario_email}  •  "
        f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}  •  "
        f"Total: {len(tareas)} tarea(s)",
        subtitulo_estilo
    ))
    elementos.append(Spacer(1, 0.5 * cm))

    if not tareas:
        elementos.append(Paragraph("No hay tareas para mostrar.", estilos["Normal"]))
    else:
        datos = [["ID", "Título", "Descripción", "Prioridad", "Fecha límite", "Estado"]]
        for t in tareas:
            if t.completada:
                estado = "Completada"
            elif _esta_vencida(t):
                estado = "VENCIDA"
            else:
                estado = "Pendiente"

            datos.append([
                str(t.id),
                t.titulo,
                (t.descripcion[:40] + "…") if t.descripcion and len(t.descripcion) > 40 else (t.descripcion or "—"),
                t.prioridad.capitalize(),
                t.fecha_limite or "—",
                estado,
            ])

        tabla = Table(datos, colWidths=[1.2 * cm, 4.5 * cm, 6 * cm, 2.5 * cm, 2.8 * cm, 2.5 * cm])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + COLOR_DARK)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCDCDC")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ]))

        # Colorear según estado
        for i, t in enumerate(tareas, start=1):
            if t.completada:
                tabla.setStyle(TableStyle([("TEXTCOLOR", (5, i), (5, i), colors.HexColor("#" + COLOR_GRAY))]))
            elif _esta_vencida(t):
                tabla.setStyle(TableStyle([
                    ("TEXTCOLOR", (5, i), (5, i), colors.HexColor("#" + COLOR_ALTA)),
                    ("FONTNAME", (5, i), (5, i), "Helvetica-Bold"),
                ]))

        elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)
    return buffer.getvalue()


def generar_pdf_tarea(tarea, usuario_email):
    """Genera un PDF con una sola tarea. Devuelve bytes."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    estilos = getSampleStyleSheet()
    titulo_estilo = ParagraphStyle(
        "TituloPersonalizado",
        parent=estilos["Title"],
        textColor=colors.HexColor("#" + COLOR_DARK),
        fontSize=18,
    )
    label_estilo = ParagraphStyle(
        "Label",
        parent=estilos["Normal"],
        textColor=colors.HexColor("#" + COLOR_GRAY),
        fontSize=9,
    )
    valor_estilo = ParagraphStyle(
        "Valor",
        parent=estilos["Normal"],
        textColor=colors.HexColor("#" + COLOR_DARK),
        fontSize=11,
    )

    elementos = []
    elementos.append(Paragraph(f"📋 Tarea #{tarea.id}", titulo_estilo))
    elementos.append(Paragraph(
        f"Usuario: {usuario_email}  •  Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}",
        label_estilo
    ))
    elementos.append(Spacer(1, 0.8 * cm))

    def fila(label, valor):
        elementos.append(Paragraph(label.upper(), label_estilo))
        elementos.append(Paragraph(str(valor), valor_estilo))
        elementos.append(Spacer(1, 0.4 * cm))

    fila("Título", tarea.titulo)
    fila("Descripción", tarea.descripcion or "—")
    fila("Prioridad", tarea.prioridad.capitalize())
    fila("Fecha límite", tarea.fecha_limite or "—")
    fila("Estado", "✅ Completada" if tarea.completada else ("⚠️ Vencida" if _esta_vencida(tarea) else "⏳ Pendiente"))
    fila("Fecha de creación", tarea.fecha_creacion.strftime("%d/%m/%Y a las %H:%M") if tarea.fecha_creacion else "—")

    doc.build(elementos)
    buffer.seek(0)
    return buffer.getvalue()


# ====================================================
# EXCEL
# ====================================================
def generar_excel_tareas(tareas, usuario_email):
    """Genera un Excel con todas las tareas. Devuelve bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tareas"

    # Título
    ws.merge_cells("A1:F1")
    ws["A1"] = "📋 Mis Tareas"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=COLOR_DARK)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Usuario: {usuario_email}  •  Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}  •  Total: {len(tareas)} tarea(s)"
    ws["A2"].font = Font(size=9, italic=True, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Encabezados
    encabezados = ["ID", "Título", "Descripción", "Prioridad", "Fecha límite", "Estado"]
    for col, enc in enumerate(encabezados, start=1):
        celda = ws.cell(row=4, column=col, value=enc)
        celda.font = Font(bold=True, color="FFFFFF", size=11)
        celda.fill = PatternFill("solid", fgColor=COLOR_DARK)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 22

    borde = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    fill_prioridad = {
        "alta": PatternFill("solid", fgColor="FDEAEA"),
        "media": PatternFill("solid", fgColor="FFF8E7"),
        "baja": PatternFill("solid", fgColor="EAFAF1"),
    }

    fila = 5
    for t in tareas:
        if t.completada:
            estado, color_estado = "✅ Completada", COLOR_GRAY
        elif _esta_vencida(t):
            estado, color_estado = "⚠️ Vencida", COLOR_ALTA
        else:
            estado, color_estado = "⏳ Pendiente", "2980B9"

        valores = [t.id, t.titulo, t.descripcion or "", t.prioridad.capitalize(),
                   t.fecha_limite or "—", estado]
        for col, val in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=val)
            celda.border = borde
            celda.alignment = Alignment(vertical="center",
                                        horizontal="left" if col in (2, 3) else "center")
            celda.fill = fill_prioridad.get(t.prioridad, PatternFill())
        ws.cell(row=fila, column=6).font = Font(bold=True, color=color_estado)
        ws.row_dimensions[fila].height = 20
        fila += 1

    for col_letra, ancho in {"A": 6, "B": 30, "C": 40, "D": 12, "E": 15, "F": 18}.items():
        ws.column_dimensions[col_letra].width = ancho
    ws.freeze_panes = "A5"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generar_excel_tarea(tarea, usuario_email):
    """Genera un Excel con una sola tarea. Devuelve bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Tarea_{tarea.id}"

    ws.merge_cells("A1:B1")
    ws["A1"] = f"📋 Tarea #{tarea.id}"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=COLOR_DARK)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    campos = [
        ("Título", tarea.titulo),
        ("Descripción", tarea.descripcion or "—"),
        ("Prioridad", tarea.prioridad.capitalize()),
        ("Fecha límite", tarea.fecha_limite or "—"),
        ("Estado", "✅ Completada" if tarea.completada else ("⚠️ Vencida" if _esta_vencida(tarea) else "⏳ Pendiente")),
        ("Fecha de creación", tarea.fecha_creacion.strftime("%d/%m/%Y a las %H:%M") if tarea.fecha_creacion else "—"),
        ("Usuario", usuario_email),
    ]

    borde = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    fila = 3
    for label, valor in campos:
        c1 = ws.cell(row=fila, column=1, value=label)
        c1.font = Font(bold=True, color=COLOR_DARK)
        c1.fill = PatternFill("solid", fgColor="ECF0F1")
        c1.border = borde
        c1.alignment = Alignment(vertical="center")

        c2 = ws.cell(row=fila, column=2, value=str(valor))
        c2.border = borde
        c2.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[fila].height = 22
        fila += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ====================================================
# TXT
# ====================================================
def generar_txt_tarea(tarea, usuario_email):
    """Genera un TXT con una sola tarea. Devuelve string."""
    estado = "✅ Completada" if tarea.completada else ("⚠️ Vencida" if _esta_vencida(tarea) else "⏳ Pendiente")

    lineas = [
        "=" * 50,
        f"📋 TAREA #{tarea.id}",
        "=" * 50,
        "",
        f"Título:          {tarea.titulo}",
        f"Descripción:     {tarea.descripcion or '—'}",
        f"Prioridad:       {tarea.prioridad.capitalize()}",
        f"Fecha límite:    {tarea.fecha_limite or '—'}",
        f"Estado:          {estado}",
        f"Fecha creación:  {tarea.fecha_creacion.strftime('%d/%m/%Y a las %H:%M') if tarea.fecha_creacion else '—'}",
        "",
        "=" * 50,
        f"Generado por:    {usuario_email}",
        f"Fecha:           {datetime.now().strftime('%d/%m/%Y a las %H:%M')}",
        "=" * 50,
    ]
    return "\n".join(lineas)


# ====================================================
# UTILIDADES
# ====================================================
def _esta_vencida(tarea):
    """Devuelve True si la tarea tiene fecha límite pasada y no está completada."""
    if not tarea.fecha_limite or tarea.completada:
        return False
    try:
        fecha = datetime.strptime(tarea.fecha_limite, "%Y-%m-%d").date()
        return datetime.now().date() > fecha
    except ValueError:
        return False


def sanitizar_nombre(texto):
    """Convierte un título en un nombre de archivo seguro."""
    import re
    texto = re.sub(r"[^\w\s-]", "", texto).strip()
    texto = re.sub(r"[-\s]+", "_", texto)
    return texto[:50] or "tarea"